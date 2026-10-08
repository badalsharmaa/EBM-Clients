import json
import re
from urllib.parse import quote
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

def create_crm_workbook():
    with open('clients.json', 'r', encoding='utf-8') as f:
        clients = json.load(f)

    wb = Workbook()
    
    # ----------------------------------------------------
    # Styles Definition
    # ----------------------------------------------------
    primary_color = "1E3A8A"     # Deep Blue
    primary_light = "DBEAFE"     # Light Blue
    header_fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    sub_header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    summary_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    card_fill = PatternFill(start_color="F0FDFA", end_color="F0FDFA", fill_type="solid")
    
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="475569")
    bold_font = Font(name="Calibri", size=11, bold=True, color="0F172A")
    regular_font = Font(name="Calibri", size=11, color="0F172A")
    kpi_num_font = Font(name="Calibri", size=20, bold=True, color="1E3A8A")
    
    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom = Border(bottom=Side(border_style="medium", color="1E3A8A"))
    card_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    wrap_left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # ----------------------------------------------------
    # Sheet 1: Dashboard / Summary
    # ----------------------------------------------------
    ws_dash = wb.active
    ws_dash.title = "Executive Dashboard"
    ws_dash.views.sheetView[0].showGridLines = True

    ws_dash["B2"] = "EBM CLIENTS - EXECUTIVE CRM DASHBOARD"
    ws_dash["B2"].font = title_font
    ws_dash["B3"] = "Real-time pipeline metrics & outreach performance from 'Lead Pipeline' tab"
    ws_dash["B3"].font = subtitle_font

    # KPI Summary Cards (pointing to Column O for Status)
    kpis = [
        ("Total Leads", "=COUNTA('Lead Pipeline'!A5:A34)", "B5", "B6", "C6"),
        ("Uncontacted", "=COUNTIF('Lead Pipeline'!O5:O34, \"1. Uncontacted\")", "D5", "D6", "E6"),
        ("Contacted / In Progress", "=COUNTIF('Lead Pipeline'!O5:O34, \"2. Contacted\") + COUNTIF('Lead Pipeline'!O5:O34, \"3. Meeting Pitched\") + COUNTIF('Lead Pipeline'!O5:O34, \"4. Proposal Sent\")", "F5", "F6", "G6"),
        ("Won / Deals Closed", "=COUNTIF('Lead Pipeline'!O5:O34, \"5. Closed - Won\")", "H5", "H6", "I6")
    ]

    for label, formula, top_left, cell_val, cell_end in kpis:
        # Merge top row for label
        start_col = top_left[0]
        row_top = int(top_left[1])
        ws_dash.merge_cells(f"{top_left}:{chr(ord(start_col)+1)}{row_top}")
        lbl_cell = ws_dash[top_left]
        lbl_cell.value = label
        lbl_cell.font = Font(name="Calibri", size=10, bold=True, color="475569")
        lbl_cell.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        lbl_cell.alignment = center_align

        # Merge value row
        ws_dash.merge_cells(f"{cell_val}:{cell_end}")
        v_cell = ws_dash[cell_val]
        v_cell.value = formula
        v_cell.font = kpi_num_font
        v_cell.fill = card_fill
        v_cell.alignment = center_align

        # Apply borders
        for r in range(row_top, row_top+2):
            for c in range(ord(start_col)-ord('A')+1, ord(start_col)-ord('A')+3):
                ws_dash.cell(row=r, column=c).border = card_border

    # Status Breakdown Table
    ws_dash["B9"] = "Lead Status Breakdown"
    ws_dash["B9"].font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    ws_dash.merge_cells("B9:D9")

    ws_dash["B10"] = "Pipeline Stage"
    ws_dash["C10"] = "Lead Count"
    ws_dash["D10"] = "Share (%)"
    for col in ["B10", "C10", "D10"]:
        ws_dash[col].fill = header_fill
        ws_dash[col].font = header_font
        ws_dash[col].alignment = center_align
        ws_dash[col].border = cell_border

    statuses = [
        ("1. Uncontacted", "=COUNTIF('Lead Pipeline'!O$5:O$34, B11)"),
        ("2. Contacted", "=COUNTIF('Lead Pipeline'!O$5:O$34, B12)"),
        ("3. Meeting Pitched", "=COUNTIF('Lead Pipeline'!O$5:O$34, B13)"),
        ("4. Proposal Sent", "=COUNTIF('Lead Pipeline'!O$5:O$34, B14)"),
        ("5. Closed - Won", "=COUNTIF('Lead Pipeline'!O$5:O$34, B15)"),
        ("6. Closed - Lost", "=COUNTIF('Lead Pipeline'!O$5:O$34, B16)")
    ]

    for idx, (st, form) in enumerate(statuses, start=11):
        ws_dash[f"B{idx}"] = st
        ws_dash[f"B{idx}"].font = regular_font
        ws_dash[f"B{idx}"].border = cell_border

        ws_dash[f"C{idx}"] = form
        ws_dash[f"C{idx}"].font = bold_font
        ws_dash[f"C{idx}"].alignment = center_align
        ws_dash[f"C{idx}"].border = cell_border

        ws_dash[f"D{idx}"] = f"=C{idx}/$C$17"
        ws_dash[f"D{idx}"].font = regular_font
        ws_dash[f"D{idx}"].alignment = right_align
        ws_dash[f"D{idx}"].number_format = "0.0%"
        ws_dash[f"D{idx}"].border = cell_border

    # Total row
    ws_dash["B17"] = "Total"
    ws_dash["B17"].font = bold_font
    ws_dash["B17"].fill = summary_fill
    ws_dash["B17"].border = cell_border

    ws_dash["C17"] = "=SUM(C11:C16)"
    ws_dash["C17"].font = bold_font
    ws_dash["C17"].alignment = center_align
    ws_dash["C17"].fill = summary_fill
    ws_dash["C17"].border = cell_border

    ws_dash["D17"] = "=SUM(D11:D16)"
    ws_dash["D17"].font = bold_font
    ws_dash["D17"].alignment = right_align
    ws_dash["D17"].number_format = "0.0%"
    ws_dash["D17"].fill = summary_fill
    ws_dash["D17"].border = cell_border

    # Industry Category Breakdown Table
    ws_dash["F9"] = "Industry Category Breakdown"
    ws_dash["F9"].font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    ws_dash.merge_cells("F9:H9")

    ws_dash["F10"] = "Industry"
    ws_dash["G10"] = "Count"
    ws_dash["H10"] = "Share (%)"
    for col in ["F10", "G10", "H10"]:
        ws_dash[col].fill = header_fill
        ws_dash[col].font = header_font
        ws_dash[col].alignment = center_align
        ws_dash[col].border = cell_border

    industries = [
        ("Locks & Hardware", "=COUNTIF('Lead Pipeline'!E$5:E$34, F11)"),
        ("Healthcare & Hospitals", "=COUNTIF('Lead Pipeline'!E$5:E$34, F12)"),
        ("Logistics & Transport", "=COUNTIF('Lead Pipeline'!E$5:E$34, F13)"),
        ("Hospitality & Hotels", "=COUNTIF('Lead Pipeline'!E$5:E$34, F14)"),
        ("Chemical & Pharma", "=COUNTIF('Lead Pipeline'!E$5:E$34, F15)")
    ]

    for idx, (ind, form) in enumerate(industries, start=11):
        ws_dash[f"F{idx}"] = ind
        ws_dash[f"F{idx}"].font = regular_font
        ws_dash[f"F{idx}"].border = cell_border

        ws_dash[f"G{idx}"] = form
        ws_dash[f"G{idx}"].font = bold_font
        ws_dash[f"G{idx}"].alignment = center_align
        ws_dash[f"G{idx}"].border = cell_border

        ws_dash[f"H{idx}"] = f"=G{idx}/$G$16"
        ws_dash[f"H{idx}"].font = regular_font
        ws_dash[f"H{idx}"].alignment = right_align
        ws_dash[f"H{idx}"].number_format = "0.0%"
        ws_dash[f"H{idx}"].border = cell_border

    # Total row
    ws_dash["F16"] = "Total"
    ws_dash["F16"].font = bold_font
    ws_dash["F16"].fill = summary_fill
    ws_dash["F16"].border = cell_border

    ws_dash["G16"] = "=SUM(G11:G15)"
    ws_dash["G16"].font = bold_font
    ws_dash["G16"].alignment = center_align
    ws_dash["G16"].fill = summary_fill
    ws_dash["G16"].border = cell_border

    ws_dash["H16"] = "=SUM(H11:H15)"
    ws_dash["H16"].font = bold_font
    ws_dash["H16"].alignment = right_align
    ws_dash["H16"].number_format = "0.0%"
    ws_dash["H16"].fill = summary_fill
    ws_dash["H16"].border = cell_border

    # Quick Guidance note on dashboard
    ws_dash["B19"] = "HOW TO USE THIS WORKBOOK:"
    ws_dash["B19"].font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
    notes = [
        "1. Go to the 'Lead Pipeline' tab to see all 30 enriched company accounts with email, website, and addresses.",
        "2. Click the 'WhatsApp Direct' link in column R to immediately chat with decision makers.",
        "3. Click the 'Website' link in column I to inspect each company's product catalog before calling.",
        "4. Use the dropdown in 'Status' (Col O) and 'Priority' (Col P) to track your interaction stages.",
        "5. Check the 'Outreach Scripts & Guides' tab for ready-to-use Email, SMS, Call, and LinkedIn templates."
    ]
    for r_idx, note in enumerate(notes, start=20):
        ws_dash[f"B{r_idx}"] = note
        ws_dash[f"B{r_idx}"].font = Font(name="Calibri", size=10, italic=True, color="334155")

    # Column dimensions for Dashboard
    ws_dash.column_dimensions["A"].width = 3
    ws_dash.column_dimensions["B"].width = 24
    ws_dash.column_dimensions["C"].width = 14
    ws_dash.column_dimensions["D"].width = 14
    ws_dash.column_dimensions["E"].width = 4
    ws_dash.column_dimensions["F"].width = 26
    ws_dash.column_dimensions["G"].width = 12
    ws_dash.column_dimensions["H"].width = 14
    ws_dash.column_dimensions["I"].width = 4

    # ----------------------------------------------------
    # Sheet 2: Lead Pipeline
    # ----------------------------------------------------
    ws_leads = wb.create_sheet(title="Lead Pipeline")
    ws_leads.views.sheetView[0].showGridLines = True

    # Title block
    ws_leads["A1"] = "EBM CLIENTS - ACTIVE ENRICHED OUTREACH PIPELINE"
    ws_leads["A1"].font = title_font
    ws_leads["A2"] = "Master list of 30 prospective accounts with contact intelligence, verified decision makers, emails, and EBM pitches"
    ws_leads["A2"].font = subtitle_font

    headers = [
        "Rank",                   # A
        "Company Name",           # B
        "Operating Brand",        # C
        "Decision Maker",         # D
        "Industry Category",      # E
        "Primary Mobile / WA",    # F
        "Office Phone",           # G
        "Email Address",          # H
        "Official Website",       # I
        "Aligarh Address",        # J
        "Products / Offerings",   # K
        "EBM Pitch Opportunity",  # L
        "Identified Pain Point",  # M
        "Outreach Hook",          # N
        "Status",                 # O
        "Priority",               # P
        "Last Contact Date",      # Q
        "WhatsApp Direct",        # R
        "Next Follow-up",         # S
        "Notes & Objections"      # T
    ]

    header_row = 4
    for col_idx, h in enumerate(headers, start=1):
        c = ws_leads.cell(row=header_row, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center_align
        c.border = cell_border

    # Categorization helper
    def categorize(company, opp):
        c_lower = company.lower()
        o_lower = opp.lower()
        if "hospital" in c_lower or "hospital" in o_lower:
            return "Healthcare & Hospitals"
        elif "logistics" in c_lower or "transport" in c_lower or "fleet" in o_lower:
            return "Logistics & Transport"
        elif "hotel" in c_lower or "guest" in o_lower:
            return "Hospitality & Hotels"
        elif "chemical" in c_lower or "pharma" in c_lower:
            return "Chemical & Pharma"
        else:
            return "Locks & Hardware"

    # Clean phone for WhatsApp
    def clean_phone(phone_str):
        if not phone_str or "verify" in phone_str.lower() or "not surfaced" in phone_str.lower():
            return None
        first_num = phone_str.split('/')[0]
        digits = re.sub(r'\D', '', first_num)
        if len(digits) == 10:
            return f"91{digits}"
        elif len(digits) > 10 and digits.startswith("91"):
            return digits
        elif len(digits) > 10:
            return f"91{digits[-10:]}"
        return None

    # Insert client rows
    for r_idx, c_data in enumerate(clients, start=5):
        rank = c_data.get("rank", r_idx - 4)
        comp = c_data.get("company", "").strip()
        brand = c_data.get("brand_name", comp)
        dm = c_data.get("decision_maker") or "Not publicly verified"
        phone = c_data.get("contact", "")
        office = c_data.get("office_phone", "")
        email = c_data.get("email", "")
        website = c_data.get("website", "")
        addr = c_data.get("address", "")
        prods = c_data.get("products", "")
        opp = c_data.get("ebm_opportunity", "")
        pain = c_data.get("pain_point", "")
        hook = c_data.get("outreach_hook", "")
        cat = categorize(comp, opp)
        
        wa_number = clean_phone(phone)
        
        ws_leads.cell(row=r_idx, column=1, value=rank).alignment = center_align
        ws_leads.cell(row=r_idx, column=2, value=comp).alignment = wrap_left_align
        ws_leads.cell(row=r_idx, column=3, value=brand).alignment = wrap_left_align
        ws_leads.cell(row=r_idx, column=4, value=dm).alignment = wrap_left_align
        ws_leads.cell(row=r_idx, column=5, value=cat).alignment = center_align
        ws_leads.cell(row=r_idx, column=6, value=phone).alignment = center_align
        ws_leads.cell(row=r_idx, column=7, value=office).alignment = center_align
        
        # Email with mailto hyperlink
        cell_email = ws_leads.cell(row=r_idx, column=8)
        if email:
            cell_email.value = f'=HYPERLINK("mailto:{email}", "{email}")'
            cell_email.font = Font(name="Calibri", size=10, color="2563EB", underline="single")
        else:
            cell_email.value = ""
        cell_email.alignment = wrap_left_align

        # Website with hyperlink
        cell_web = ws_leads.cell(row=r_idx, column=9)
        if website:
            cell_web.value = f'=HYPERLINK("{website}", "{website.replace("https://","").replace("http://","").replace("www.","").split("/")[0]}")'
            cell_web.font = Font(name="Calibri", size=10, color="2563EB", underline="single")
        else:
            cell_web.value = ""
        cell_web.alignment = center_align

        ws_leads.cell(row=r_idx, column=10, value=addr).alignment = wrap_left_align
        ws_leads.cell(row=r_idx, column=11, value=prods).alignment = wrap_left_align
        ws_leads.cell(row=r_idx, column=12, value=opp).alignment = wrap_left_align
        ws_leads.cell(row=r_idx, column=13, value=pain).alignment = wrap_left_align
        ws_leads.cell(row=r_idx, column=14, value=hook).alignment = wrap_left_align

        # Default status & priority
        ws_leads.cell(row=r_idx, column=15, value="1. Uncontacted").alignment = center_align
        default_pri = "High" if rank <= 10 else "Medium"
        ws_leads.cell(row=r_idx, column=16, value=default_pri).alignment = center_align
        
        ws_leads.cell(row=r_idx, column=17, value="").alignment = center_align # Last contact date
        
        # WhatsApp formula
        cell_wa = ws_leads.cell(row=r_idx, column=18)
        if wa_number:
            recip = dm.split('/')[0].split('(')[0].split('—')[0].strip() if not dm.lower().startswith('not') else 'Sir/Ma\'am'
            msg = f"Namaste {recip}, Badal here from EBM Solutions Aligarh. Regarding {comp}: we specialize in {opp}. Would you be open for a quick 5-min discussion this week? (Reply STOP to opt out)"
            enc_msg = quote(msg)
            wa_url = f"https://wa.me/{wa_number}?text={enc_msg}"
            cell_wa.value = f'=HYPERLINK("{wa_url}", "💬 Chat on WhatsApp")'
            cell_wa.font = Font(name="Calibri", size=10, color="047857", underline="single", bold=True)
        else:
            cell_wa.value = "Verify Number"
            cell_wa.font = Font(name="Calibri", size=10, italic=True, color="94A3B8")
        cell_wa.alignment = center_align

        ws_leads.cell(row=r_idx, column=19, value="").alignment = center_align # Next follow-up
        ws_leads.cell(row=r_idx, column=20, value="").alignment = wrap_left_align # Notes

        # Row styling
        fill_to_use = zebra_fill if (r_idx % 2 == 0) else PatternFill(fill_type=None)
        for c in range(1, 21):
            cell = ws_leads.cell(row=r_idx, column=c)
            cell.border = cell_border
            if cell.font.name != "Calibri" or (cell.font.color and cell.font.color.rgb not in ["047857", "94A3B8", "2563EB"]):
                cell.font = regular_font
            if fill_to_use.fill_type:
                cell.fill = fill_to_use

    # Add Dropdown Validation for Status and Priority
    status_dv = DataValidation(type="list", formula1='"1. Uncontacted,2. Contacted,3. Meeting Pitched,4. Proposal Sent,5. Closed - Won,6. Closed - Lost"', allow_blank=True)
    status_dv.error = 'Your entry is not in the list'
    status_dv.errorTitle = 'Invalid Status'
    status_dv.prompt = 'Select status from the list'
    status_dv.promptTitle = 'Lead Status'
    ws_leads.add_data_validation(status_dv)
    status_dv.add(f"O5:O{len(clients)+4}")

    priority_dv = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True)
    priority_dv.error = 'Select High, Medium, or Low'
    priority_dv.errorTitle = 'Invalid Priority'
    ws_leads.add_data_validation(priority_dv)
    priority_dv.add(f"P5:P{len(clients)+4}")

    # Column widths for Lead Pipeline
    col_widths = {
        "A": 8,   # Rank
        "B": 30,  # Company Name
        "C": 26,  # Brand
        "D": 26,  # Decision Maker
        "E": 22,  # Industry
        "F": 20,  # Primary Mobile / WA
        "G": 20,  # Office Phone
        "H": 26,  # Email
        "I": 22,  # Website
        "J": 35,  # Address
        "K": 35,  # Products
        "L": 35,  # EBM Pitch
        "M": 35,  # Pain Point
        "N": 35,  # Outreach Hook
        "O": 18,  # Status
        "P": 12,  # Priority
        "Q": 16,  # Last Contact
        "R": 22,  # WhatsApp Link
        "S": 16,  # Next Follow-up
        "T": 32   # Notes
    }
    for col_letter, width in col_widths.items():
        ws_leads.column_dimensions[col_letter].width = width

    # Enable autofilter on lead table
    ws_leads.auto_filter.ref = f"A4:T{len(clients)+4}"

    # ----------------------------------------------------
    # Sheet 3: Outreach Scripts & Playbook
    # ----------------------------------------------------
    ws_scripts = wb.create_sheet(title="Outreach Scripts & Guides")
    ws_scripts.views.sheetView[0].showGridLines = True

    ws_scripts["B2"] = "COLD OUTREACH SCRIPTS & ACTION PLAYBOOK"
    ws_scripts["B2"].font = title_font
    ws_scripts["B3"] = "Direct templates distilled from deep-research-report.md"
    ws_scripts["B3"].font = subtitle_font

    scripts_data = [
        ("WhatsApp / SMS Quick Pitch", 
         "Hi [Name], Badal here from EBM Solutions in Aligarh. We help [Industry/Locks/Hospitals] in Aligarh automate their dealer inquiries and customer management. We built a solution that saved similar businesses 10+ hours a week and boosted inquiries by 40%. Could we do a quick 5-min call this week? Reply Y/N. (P.S. Reply STOP to opt-out)"),
        
        ("Cold Email Pitch",
         "Subject: Modernizing [Company Name]'s dealer & sales operations\n\nHi [Decision Maker],\n\nI noticed [Company Name]'s strong reputation in Aligarh. We recently worked with peer firms to implement [ebm_opportunity], which cut inquiry response times and drove qualified distributor leads.\n\nKey advantage for [Company Name]:\n• [outreach_hook]\n\nWould you be open to a 10-minute introductory conversation this Thursday or Friday?\n\nBest regards,\nBadal Sharma | EBM Solutions\nPhone: [Your Phone]\n(Reply STOP to unsubscribe)"),
        
        ("LinkedIn Connection Note",
         "Hi [Name], I'm working with manufacturing and healthcare leaders in Uttar Pradesh to modernize their sales and dealer workflows. Would love to connect and share brief insights into digital catalog & CRM automation!"),
        
        ("Phone Call Script (Opening 30 Seconds)",
         "\"Namaste [Name / Manager], this is Badal from EBM Solutions in Aligarh. I know you're busy so I'll be brief. We specifically assist [Industry, e.g. lock manufacturers / hospitals] in Aligarh with custom digital systems for [ebm_opportunity]. I wanted to see who on your team handles operations and sales systems?\"")
    ]

    cur_row = 6
    for title, script in scripts_data:
        ws_scripts[f"B{cur_row}"] = title
        ws_scripts[f"B{cur_row}"].font = Font(name="Calibri", size=12, bold=True, color="1E3A8A")
        ws_scripts[f"B{cur_row}"].fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
        ws_scripts.merge_cells(f"B{cur_row}:F{cur_row}")
        
        cur_row += 1
        ws_scripts[f"B{cur_row}"] = script
        ws_scripts[f"B{cur_row}"].font = regular_font
        ws_scripts[f"B{cur_row}"].alignment = wrap_left_align
        ws_scripts.merge_cells(f"B{cur_row}:F{cur_row+2}")
        
        for r in range(cur_row-1, cur_row+3):
            for c in range(2, 7):
                ws_scripts.cell(row=r, column=c).border = cell_border
                
        cur_row += 4

    ws_scripts.column_dimensions["A"].width = 3
    ws_scripts.column_dimensions["B"].width = 25
    ws_scripts.column_dimensions["C"].width = 25
    ws_scripts.column_dimensions["D"].width = 25
    ws_scripts.column_dimensions["E"].width = 25
    ws_scripts.column_dimensions["F"].width = 25

    # ----------------------------------------------------
    # Sheet 4: 30-60-90 Day Action Roadmap
    # ----------------------------------------------------
    ws_plan = wb.create_sheet(title="30-60-90 Day Plan")
    ws_plan.views.sheetView[0].showGridLines = True

    ws_plan["B2"] = "30 / 60 / 90 DAY EXECUTION ROADMAP"
    ws_plan["B2"].font = title_font
    ws_plan["B3"] = "Phased timeline for data enrichment, multi-channel outreach, and closing deals"
    ws_plan["B3"].font = subtitle_font

    plan_headers = ["Phase", "Timeframe", "Core Objective", "Key Action Items", "Success Target (KPI)"]
    for col_idx, h in enumerate(plan_headers, start=2):
        c = ws_plan.cell(row=5, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center_align
        c.border = cell_border

    plan_data = [
        ("Phase 1: Foundation & Audit", "Month 1 (Days 1–30)", "Data Enrichment & CRM Setup", 
         "• Audit all 30 pre-loaded leads in this sheet\n• Fill in decision maker names, verified emails & LinkedIn profiles\n• Test WhatsApp & email outreach scripts with top 5 lock manufacturers\n• Ensure opt-out mechanics (STOP / Unsubscribe) are active", 
         "100% records enriched with direct contacts; 15 initial conversations initiated"),
        
        ("Phase 2: Active Multi-Channel Outreach", "Month 2 (Days 31–60)", "Pipeline Velocity & Pitches",
         "• Roll out WhatsApp & email sequences to all 30 target leads\n• Launch local Google Business Profile & local SEO for EBM\n• Conduct 8–10 discovery & demo calls\n• Send formal proposals for Dealer CRM / Custom Automation", 
         "At least 6 formal proposals sent; 2 deals in final negotiation"),
        
        ("Phase 3: Deal Closing & Expansion", "Month 3 (Days 61–90)", "Closing & Referral Engine",
         "• Close initial 2–3 pilot accounts (Offer introductory retainer/pilot discount)\n• Implement client solutions & measure time/cost savings\n• Request testimonials and referrals to fellow Aligarh industrial units\n• Expand pipeline to 30 additional regional accounts", 
         "2–3 paying clients acquired; ₹1.5L+ in contracted revenue; 1 published case study")
    ]

    for idx, (phase, tf, obj, actions, kpi) in enumerate(plan_data, start=6):
        ws_plan.cell(row=idx, column=2, value=phase).font = bold_font
        ws_plan.cell(row=idx, column=2).alignment = left_align
        
        ws_plan.cell(row=idx, column=3, value=tf).font = regular_font
        ws_plan.cell(row=idx, column=3).alignment = center_align
        
        ws_plan.cell(row=idx, column=4, value=obj).font = bold_font
        ws_plan.cell(row=idx, column=4).alignment = left_align
        
        ws_plan.cell(row=idx, column=5, value=actions).font = regular_font
        ws_plan.cell(row=idx, column=5).alignment = wrap_left_align
        
        ws_plan.cell(row=idx, column=6, value=kpi).font = regular_font
        ws_plan.cell(row=idx, column=6).alignment = wrap_left_align
        
        fill_to_use = zebra_fill if (idx % 2 == 0) else PatternFill(fill_type=None)
        for c in range(2, 7):
            cell = ws_plan.cell(row=idx, column=c)
            cell.border = cell_border
            if fill_to_use.fill_type:
                cell.fill = fill_to_use
        ws_plan.row_dimensions[idx].height = 65

    ws_plan.column_dimensions["A"].width = 3
    ws_plan.column_dimensions["B"].width = 28
    ws_plan.column_dimensions["C"].width = 20
    ws_plan.column_dimensions["D"].width = 28
    ws_plan.column_dimensions["E"].width = 46
    ws_plan.column_dimensions["F"].width = 34

    wb.save('EBM_Clients_Lead_Management_CRM.xlsx')
    print("Successfully created 'EBM_Clients_Lead_Management_CRM.xlsx' with 20 enriched columns across 30 leads.")

if __name__ == '__main__':
    create_crm_workbook()
